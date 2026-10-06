"""ชุดทดสอบทั้งระบบ: รันด้วย  pytest -q
ใช้ผู้ให้บริการ AI จำลอง (ไม่ออกอินเทอร์เน็ต ไม่ใช้โควตา) เพื่อทดสอบเส้นทางเรียก Gemini และ OpenAI-compatible จริง
"""
import io
import json
import os
import random
import tempfile
import zipfile

import httpx
import pytest
from PIL import Image

_tmp = tempfile.mkdtemp(prefix="fives_test_")
os.environ.setdefault("DATA_DIR", _tmp)
os.environ["DISABLE_WORKER"] = "1"
os.environ["ADMIN_PASSWORD"] = "admin1234"

from fastapi.testclient import TestClient  # noqa: E402

from app import ai, db as dbm, settings_store, storage, worker  # noqa: E402
from app.main import app  # noqa: E402

AI = {"gemini": "ok", "openai": "ok", "levels": [4, 3, 2, 3, 4], "calls": [], "wrap": False, "image_ok": True, "na": False}


def answer():
    codes = ["S1", "S2", "S3", "S4", "S5"]
    crit = [dict(code=c, na=False, level=AI["levels"][i], reason=f"เหตุผลของ {c}", findings=[f"พบที่ {c}"],
                 recommendations=[f"แก้ไข {c}"]) for i, c in enumerate(codes)]
    if AI["na"]:
        crit[4] = dict(code="S5", na=True, level=None, reason="ไม่เห็นหลักฐาน")
        crit[0]["na"] = True       # S1 ไม่อนุญาตให้ข้าม ระบบต้องไม่ยอม
    text = json.dumps(dict(image_ok=AI["image_ok"], image_issue="" if AI["image_ok"] else "ภาพเบลอ", scene="ชั้นวางของ",
                           criteria=crit, summary="สรุปผล", top_actions=["ทำข้อ 1", "ทำข้อ 2"],
                           overview="ภาพรวมแผนก", strengths=["สะอาด"],
                           priorities=[dict(issue="ของวางบนพื้น", action="ทำที่วาง", where="มุมซ้าย")]), ensure_ascii=False)
    return f"```json\n{text}\n```" if AI["wrap"] else text


def handler(request: httpx.Request) -> httpx.Response:
    url = str(request.url)
    AI["calls"].append(url)
    if "/v1beta/models" in url:
        assert request.headers.get("x-goog-api-key") == "g-key"
        if request.method == "GET":
            return httpx.Response(200, json={"models": [
                {"name": "models/gemini-3.8-flash", "supportedGenerationMethods": ["generateContent"]},
                {"name": "models/gemini-embedding-001", "supportedGenerationMethods": ["embedContent"]}]})
        if AI["gemini"] == "429":
            return httpx.Response(429, json={"error": {"message": "Quota exceeded. Please retry in 34.5s."}})
        if AI["gemini"] == "403":
            return httpx.Response(403, json={"error": {"message": "API key not valid"}})
        if AI["gemini"] == "garbage":
            return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": "ขออภัย"}]}}]})
        body = json.loads(request.content)
        assert body["generationConfig"]["temperature"] == 0
        prompt = json.dumps(body, ensure_ascii=False)
        assert "SECRET-DEPT-NAME" not in prompt, "ชื่อแผนกต้องไม่ถูกส่งให้ AI"
        return httpx.Response(200, json={"candidates": [{"content": {"parts": [{"text": answer()}]}, "finishReason": "STOP"}]})
    if url.endswith("/models"):
        return httpx.Response(200, json={"data": [{"id": "vision-model-a"}]})
    if url.endswith("/chat/completions"):
        assert request.headers.get("authorization") == "Bearer o-key"
        body = json.loads(request.content)
        if AI["openai"] == "nojsonmode" and "response_format" in body:
            return httpx.Response(400, json={"error": {"message": "response_format is not supported"}})
        return httpx.Response(200, json={"choices": [{"message": {"content": answer()}}]})
    return httpx.Response(404, json={"error": "not found"})


ai._transport = httpx.MockTransport(handler)


def jpeg(seed: int, size=(1600, 1200)) -> bytes:
    rnd = random.Random(seed)
    im = Image.new("RGB", size, (rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)))
    for _ in range(30):
        x, y = rnd.randrange(size[0]), rnd.randrange(size[1])
        im.paste((rnd.randrange(256), rnd.randrange(256), rnd.randrange(256)), (x, y, x + 60, y + 40))
    buf = io.BytesIO()
    im.save(buf, "JPEG", quality=90)
    return buf.getvalue()


def login(username, password):
    c = TestClient(app, follow_redirects=False)
    r = c.post("/login", data={"username": username, "password": password})
    assert r.status_code == 303, r.text
    return c


def upload(c, rid, did, seed, area="ชั้นวาง A", expect=200):
    r = c.post("/api/photos", data=dict(round_id=rid, department_id=did, area_name=area, area_type="warehouse", note="หมายเหตุ"),
               files={"file": ("p.jpg", jpeg(seed), "image/jpeg")})
    assert r.status_code == expect, r.text
    return r.json().get("id") if expect == 200 else r.json()


def drain(limit=200):
    n = 0
    while worker.process_one() and n < limit:
        n += 1
    return n


def photo(pid):
    with dbm.SessionLocal() as s:
        return s.get(dbm.Photo, pid)


S = {}


def test_01_login_and_forced_password_change():
    c = TestClient(app, follow_redirects=False)
    assert c.get("/").headers["location"].startswith("/login")
    assert c.get("/healthz").json()["ok"] is True
    r = c.post("/login", data={"username": "admin", "password": "wrong-pass"})
    assert r.status_code == 303 and r.headers["location"] == "/login"
    c = login("admin", "admin1234")
    assert c.get("/").headers["location"] == "/account/password"      # ต้องตั้งรหัสใหม่ก่อน
    assert c.get("/admin").headers["location"] == "/account/password"
    r = c.post("/account/password", data={"current": "admin1234", "new": "short", "confirm": "short"})
    assert r.headers["location"] == "/account/password"
    r = c.post("/account/password", data={"current": "admin1234", "new": "Factory5S2026", "confirm": "Factory5S2026"})
    assert r.headers["location"] == "/"
    assert c.get("/").status_code == 200
    assert login("admin", "Factory5S2026").get("/admin").status_code == 200
    S["admin"] = c


def test_02_cross_site_post_is_blocked():
    c = S["admin"]
    r = c.post("/admin/departments/save", data={"code": "X", "name": "X"}, headers={"origin": "https://evil.example"})
    assert r.status_code == 403
    with dbm.SessionLocal() as s:
        assert s.query(dbm.Department).count() == 0


def test_03_admin_setup():
    c = S["admin"]
    for code, name in (("PRD", "SECRET-DEPT-NAME ผลิต"), ("WH", "คลังสินค้า"), ("QC", "ควบคุมคุณภาพ"), ("MT", "ซ่อมบำรุง")):
        assert c.post("/admin/departments/save", data={"code": code, "name": name}).status_code == 303
    assert c.post("/admin/departments/save", data={"code": "prd", "name": "ซ้ำ"}).status_code == 303
    with dbm.SessionLocal() as s:
        depts = {d.code: d.id for d in s.query(dbm.Department).all()}
    assert set(depts) == {"PRD", "WH", "QC", "MT"}
    S["depts"] = depts
    users = [("somchai", "member", depts["PRD"]), ("malee", "member", depts["WH"]), ("auditor1", "auditor", "")]
    for u, role, d in users:
        r = c.post("/admin/users/save", data={"username": u, "full_name": u, "role": role, "department_id": d, "password": "Start1234"})
        assert r.status_code == 303
    r = c.post("/admin/users/save", data={"username": "weak", "role": "auditor", "password": "123"})   # รหัสอ่อน ไม่ผ่าน
    with dbm.SessionLocal() as s:
        assert s.query(dbm.User).filter_by(username="weak").count() == 0
        assert s.query(dbm.User).count() == 4
    # ผู้ดูแลคนสุดท้ายลดสิทธิ์ตัวเองไม่ได้
    c.post("/admin/users/save", data={"id": 1, "username": "admin", "role": "member", "department_id": depts["PRD"]})
    with dbm.SessionLocal() as s:
        assert s.get(dbm.User, 1).role == "admin"
    # แก้เกณฑ์: คะแนนเต็ม S1 เป็น 40 -> รวม 120
    with dbm.SessionLocal() as s:
        c1 = s.query(dbm.Criterion).filter_by(code="S1").one()
        data = {"id": c1.id, "code": "S1", "name": c1.name, "max_score": "40", "focus": c1.focus, "active": "1", "sort_order": 10}
        data.update({f"level{i}": c1.levels[i] for i in range(5)})
    assert c.post("/admin/criteria/save", data=data).status_code == 303
    assert "120" in c.get("/admin/criteria").text
    r = c.post("/admin/settings", data={"org_name": "AHP ทดสอบ", "ai1_type": "gemini", "ai1_key": "g-key", "ai1_model": "gemini-3.8-flash",
                                        "ai2_type": "none", "ai_rpm": 30, "ai_daily": 500, "ai_max_attempts": 3, "img_max_side": 1280,
                                        "img_quality": 78, "allow_gallery": "1", "storage_budget_mb": 350, "retention_days": 90,
                                        "purge_requires_backup": "1", "ranking_visibility": "always", "member_see_all": "1"})
    assert r.status_code == 303
    s_ = settings_store.load()
    assert s_["ai1_key"] == "g-key" and s_["org_name"] == "AHP ทดสอบ"
    c.post("/admin/settings", data={"org_name": "AHP ทดสอบ", "ai1_type": "gemini", "ai1_key": "", "ai1_model": "gemini-3.8-flash",
                                    "ai_rpm": 30, "ai_daily": 500, "ai_max_attempts": 3, "member_see_all": "1", "allow_gallery": "1",
                                    "purge_requires_backup": "1"})
    assert settings_store.load()["ai1_key"] == "g-key"        # เว้นว่าง = ใช้ key เดิม
    assert "g-key" not in c.get("/admin/settings").text       # ไม่แสดง key เต็มในหน้าเว็บ
    r = c.post("/admin/rounds/save", data={"name": "ตรวจ 5ส ตุลาคม", "start_date": "2026-10-01", "end_date": "2026-10-31", "min_photos": 2})
    assert r.status_code == 303
    with dbm.SessionLocal() as s:
        rnd = s.query(dbm.Round).one()
        assert sum(x["max"] for x in rnd.rubric) == 120 and len(rnd.rubric) == 5
        S["rid"] = rnd.id


def test_04_ai_settings_endpoints():
    c = S["admin"]
    r = c.post("/admin/ai/models", json={"slot": "ai1", "type": "gemini", "base": "", "key": "", "model": ""})
    assert r.json() == {"ok": True, "models": ["gemini-3.8-flash"]}
    r = c.post("/admin/ai/test", json={"slot": "ai1", "type": "gemini", "base": "", "key": "", "model": "gemini-3.8-flash"})
    assert r.json()["ok"] is True and "5 เกณฑ์" in r.json()["message"]
    AI["gemini"] = "403"
    r = c.post("/admin/ai/test", json={"slot": "ai1", "type": "gemini", "key": "", "model": "gemini-3.8-flash"})
    assert r.json()["ok"] is False and "API key" in r.json()["error"]
    AI["gemini"] = "ok"
    r = c.post("/admin/ai/models", json={"slot": "ai2", "type": "openai", "base": "https://x.test/v1", "key": "o-key"})
    assert r.json()["models"] == ["vision-model-a"]


def test_05_upload_rules():
    rid, d = S["rid"], S["depts"]
    m = login("somchai", "Start1234")
    assert m.get("/capture").headers["location"] == "/account/password"
    m.post("/account/password", data={"current": "Start1234", "new": "Somchai5S99", "confirm": "Somchai5S99"})
    page = m.get("/capture")
    assert page.status_code == 200 and "คลังสินค้า</option>" not in page.text.split('id="dept"')[1].split("</select>")[0]
    S["member"] = m
    S["p1"] = upload(m, rid, d["PRD"], 1)
    assert upload(m, rid, d["WH"], 2, expect=403)["detail"]                    # ส่งแทนแผนกอื่นไม่ได้
    assert "ส่งแล้ว" in upload(m, rid, d["PRD"], 1, expect=409)["detail"]        # ภาพซ้ำ
    assert upload(m, rid, d["PRD"], 3, area="  ", expect=400)["detail"]          # ไม่ใส่จุดตรวจ
    r = m.post("/api/photos", data=dict(round_id=rid, department_id=d["PRD"], area_name="x"), files={"file": ("a.jpg", b"not an image", "image/jpeg")})
    assert r.status_code == 400
    r = m.post("/api/photos", data=dict(round_id=rid, department_id=d["PRD"], area_name="x"),
               files={"file": ("a.jpg", jpeg(77, (120, 90)), "image/jpeg")})
    assert r.status_code == 400                                                # ภาพเล็กเกินไป
    p = photo(S["p1"])
    assert p.status == "pending" and p.width == 1280 and p.image_bytes < 400_000 and p.thumb_bytes < 40_000
    img = m.get(f"/photos/{S['p1']}/image")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"
    assert Image.open(io.BytesIO(img.content)).getexif().get(0x8825) is None    # ไม่มีพิกัด GPS ติดไป
    assert TestClient(app, follow_redirects=False).get(f"/photos/{S['p1']}/image").status_code == 303


def test_06_ai_scoring_paths():
    rid, d, m = S["rid"], S["depts"], S["member"]
    assert drain() == 1
    p = photo(S["p1"])
    # ระดับ 4,3,2,3,4 กับคะแนนเต็ม 40,20,20,20,20 -> 40+15+10+15+20 = 100 จาก 120
    assert (p.status, p.score, p.max_score, round(p.percent, 2)) == ("done", 100.0, 120.0, 83.33)
    assert p.model == "gemini-3.8-flash" and p.analysis["criteria"][0]["reason"]
    # คำตอบมี ```json ครอบ + เกณฑ์ที่ข้ามได้ (S5) และข้ามไม่ได้ (S1)
    AI.update(wrap=True, na=True, levels=[2, 2, 2, 2, 0])
    p2 = upload(m, rid, d["PRD"], 10, area="โต๊ะประกอบ")
    drain()
    p = photo(p2)
    crit = {c["code"]: c for c in p.analysis["criteria"]}
    assert p.status == "done" and crit["S5"]["na"] is True and crit["S1"]["na"] is False
    assert (p.score, p.max_score, p.percent) == (50.0, 100.0, 50.0)            # S5 ไม่ถูกนับทั้งคะแนนและคะแนนเต็ม
    AI.update(wrap=False, na=False)
    # ภาพใช้ประเมินไม่ได้
    AI["image_ok"] = False
    p3 = upload(m, rid, d["PRD"], 11, area="ภาพเบลอ")
    drain()
    assert photo(p3).status == "rejected" and photo(p3).percent is None
    AI["image_ok"] = True
    # เกินโควตา -> กลับเข้าคิว รอเวลาที่ผู้ให้บริการบอก
    AI["gemini"] = "429"
    p4 = upload(m, rid, d["PRD"], 12, area="ทางเดิน")
    worker.process_one()
    p = photo(p4)
    assert p.status == "pending" and p.attempts == 1 and p.next_try_at is not None and "429" in p.error
    assert (p.next_try_at - dbm.now()).total_seconds() > 25
    assert worker.process_one() is False                                       # ยังไม่ถึงเวลา ไม่เรียกซ้ำ
    # ตั้ง AI สำรอง แล้วให้ทำงานแทนเมื่อตัวหลักเกินโควตา (และผู้ให้บริการสำรองไม่รองรับโหมด JSON)
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"ai2_type": "openai", "ai2_base": "https://x.test/v1", "ai2_key": "o-key", "ai2_model": "vision-model-a"})
        s.get(dbm.Photo, p4).next_try_at = None
        s.commit()
    AI.update(openai="nojsonmode", levels=[3, 3, 3, 3, 3])
    drain()
    p = photo(p4)
    assert p.status == "done" and p.provider == "openai" and p.model == "vision-model-a" and p.percent == 75.0
    # key ผิด + ไม่มีตัวสำรอง -> ผิดพลาดทันที ไม่วนลองซ้ำ
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"ai2_type": "none"})
    AI["gemini"] = "403"
    p5 = upload(m, rid, d["PRD"], 13, area="ตู้เครื่องมือ")
    drain()
    assert photo(p5).status == "error" and "API key" in photo(p5).error
    # คำตอบไม่ใช่ JSON -> ลองใหม่จนครบ 3 ครั้งแล้วจึงผิดพลาด
    AI["gemini"] = "garbage"
    S["admin"].post(f"/admin/photos/{p5}/reanalyze")
    for _ in range(3):
        with dbm.SessionLocal() as s:
            s.get(dbm.Photo, p5).next_try_at = None
            s.commit()
        worker.process_one()
    assert photo(p5).status == "error" and photo(p5).attempts == 3
    AI["gemini"] = "ok"
    S["admin"].post(f"/admin/rounds/{rid}/retry-errors")
    drain()
    assert photo(p5).status == "done"
    S.update(p2=p2, p3=p3, p4=p4, p5=p5)
    st = m.get(f"/api/photos/status?ids={S['p1']},{p3}").json()["photos"]
    assert {x["status"] for x in st} == {"done", "rejected"}


def test_07_ranking():
    rid, d = S["rid"], S["depts"]
    a = login("auditor1", "Start1234")
    a.post("/account/password", data={"current": "Start1234", "new": "Auditor5S99", "confirm": "Auditor5S99"})
    S["auditor"] = a
    AI["levels"] = [4, 4, 4, 4, 4]
    for seed in (20, 21, 22):
        upload(a, rid, d["WH"], seed, area=f"ชั้น {seed}")
    drain()
    AI["levels"] = [1, 1, 1, 1, 1]
    upload(a, rid, d["QC"], 30, area="โต๊ะ QC")      # ภาพเดียว ไม่ถึงขั้นต่ำ 2
    drain()
    with dbm.SessionLocal() as s:
        from app import scoring
        rk = scoring.round_ranking(s, s.get(dbm.Round, rid))
    assert [r["dept"].code for r in rk["ranked"]] == ["WH", "PRD"]
    assert rk["ranked"][0]["avg"] == 100.0 and rk["ranked"][0]["rank"] == 1 and rk["ranked"][1]["rank"] == 2
    prd = rk["ranked"][1]
    assert prd["scored"] == 4 and prd["rejected"] == 1 and prd["total"] == 5
    assert [r["dept"].code for r in rk["unranked"]] == ["QC"] and [r["dept"].code for r in rk["idle"]] == ["MT"]
    page = S["member"].get(f"/ranking?round={rid}")
    body = page.text.split('class="rank-list"')[1]
    assert page.status_code == 200 and body.index("คลังสินค้า") < body.index("SECRET-DEPT-NAME")
    assert S["member"].get("/").status_code == 200 and a.get("/photos").status_code == 200
    # ตั้งให้เห็นอันดับเมื่อปิดรอบ + ไม่ให้ดูแผนกอื่น
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"ranking_visibility": "closed", "member_see_all": False})
    assert "ยังไม่เปิดให้ดู" in S["member"].get(f"/ranking?round={rid}").text
    assert S["member"].get(f"/rounds/{rid}/dept/{d['WH']}").status_code == 403
    assert S["member"].get(f"/rounds/{rid}/dept/{d['PRD']}").status_code == 200
    with dbm.SessionLocal() as s:
        wh_photo = s.query(dbm.Photo).filter_by(department_id=d["WH"]).first().id
    assert S["member"].get(f"/photos/{wh_photo}").status_code == 403
    assert S["member"].get(f"/photos/{wh_photo}/image").status_code == 403
    assert S["member"].get(f"/rounds/{rid}/report").status_code == 403
    assert S["member"].get("/admin").status_code == 403
    assert S["member"].get("/admin/rounds/1/backup.zip").status_code == 403
    assert a.get(f"/ranking?round={rid}").status_code == 200 and "ยังไม่เปิดให้ดู" not in a.get(f"/ranking?round={rid}").text
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"ranking_visibility": "always", "member_see_all": True})


def test_08_override_and_summary():
    c, rid, d = S["admin"], S["rid"], S["depts"]
    pid = S["p2"]
    r = c.post(f"/admin/photos/{pid}/override", data={"level_S1": 4, "level_S2": 4, "level_S3": 4, "level_S4": 4, "level_S5": "na", "note": "x"})
    assert photo(pid).overridden is False                           # ไม่มีเหตุผล ไม่บันทึก
    r = c.post(f"/admin/photos/{pid}/override", data={"level_S1": 4, "level_S2": 4, "level_S3": 4, "level_S4": 4, "level_S5": "na",
                                                      "note": "AI มองไม่เห็นป้ายชี้บ่งที่ติดอยู่ด้านข้าง"})
    assert r.status_code == 303
    p = photo(pid)
    assert p.overridden and p.percent == 100.0 and p.analysis["ai_original"]["percent"] == 50.0
    assert S["member"].post(f"/admin/photos/{pid}/override", data={"note": "ขอเพิ่มคะแนน"}).status_code == 403
    page = c.get(f"/photos/{pid}")
    assert "กรรมการปรับผลของภาพนี้" in page.text and "คะแนนเดิมจาก AI" in page.text
    r = c.post(f"/admin/rounds/{rid}/dept/{d['PRD']}/summarize")
    assert r.status_code == 303
    page = c.get(f"/rounds/{rid}/dept/{d['PRD']}")
    assert "ภาพรวมแผนก" in page.text and "ทำที่วาง" in page.text
    with dbm.SessionLocal() as s:
        logs = [x.action for x in s.query(dbm.AuditLog).all()]
    assert "override_score" in logs and "summarize_dept" in logs


def test_09_pages_and_exports():
    c, rid = S["admin"], S["rid"]
    for url in ("/", "/capture", "/photos", f"/photos?round={rid}&status=done", f"/photos/{S['p1']}", f"/photos/{S['p3']}",
                "/ranking", "/admin", "/admin/rounds", "/admin/criteria", "/admin/departments", "/admin/users",
                "/admin/settings", "/admin/storage", "/admin/logs", "/admin/notifications", "/admin/cameras", "/trend",
                "/account/password", f"/rounds/{rid}/report"):
        r = c.get(url)
        assert r.status_code == 200, (url, r.text[:300])
        assert "Traceback" not in r.text
    assert c.get("/nope").status_code == 404 and "ไม่พบหน้านี้" in c.get("/nope").text
    assert c.get("/photos/99999").status_code == 404
    from openpyxl import load_workbook
    assert c.get(f"/admin/rounds/{rid}/export.xlsx").headers["location"] == f"/rounds/{rid}/export/scores.xlsx"   # ลิงก์รุ่น 1.0
    r = c.get(f"/rounds/{rid}/export/scores.xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["อันดับ", "รายละเอียดภาพ", "เกณฑ์ของรอบนี้"]
    ws = wb["อันดับ"]
    assert ws["A5"].value == 1 and ws["B5"].value == "WH" and ws["D5"].value == 100
    assert wb["รายละเอียดภาพ"].max_row == 10 and len(wb["รายละเอียดภาพ"]._images) == 9
    r = c.get("/admin/backup/system.json")
    data = r.json()
    assert len(data["departments"]) == 4 and "ai1_key" not in data["settings"] and "password" not in r.text.lower()


def test_10_backup_restore_and_purge():
    c, rid = S["admin"], S["rid"]
    with dbm.SessionLocal() as s:
        # ยังไม่เคยสำรอง: นโยบายลบอัตโนมัติต้องไม่ลบอะไร แม้ภาพเก่าเกินอายุ
        s.query(dbm.Photo).update({"created_at": dbm.now() - __import__("datetime").timedelta(days=200)})
        s.commit()
        assert storage.auto_cleanup(s, settings_store.load()) == {"by_age": 0, "by_space": 0}
    r = c.get(f"/admin/rounds/{rid}/backup.zip")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    manifest = json.loads(z.read("manifest.json"))
    assert {"manifest.json", "report.html", "scores.xlsx"} <= set(names)
    assert len([n for n in names if n.startswith("images/")]) == 9 and len(manifest["photos"]) == 9
    html = z.read("report.html").decode()
    assert 'src="images/WH/' in html and "คลังสินค้า" in html
    S["zip"] = r.content
    # สำรองแล้ว -> ภาพเก่าเกิน 90 วันถูกลบภาพเต็ม แต่คะแนนและภาพย่อยังอยู่
    with dbm.SessionLocal() as s:
        res = storage.auto_cleanup(s, settings_store.load())
        assert res["by_age"] == 9
        assert s.query(dbm.PhotoImage).count() == 0 and s.query(dbm.PhotoThumb).count() == 9
    p = photo(S["p1"])
    assert p.has_image is False and p.percent is not None
    img = c.get(f"/photos/{S['p1']}/image")
    assert img.status_code == 200 and len(img.content) == p.thumb_bytes      # ได้ภาพย่อแทน
    assert "ลบภาพเต็มแล้ว" in c.get(f"/photos/{S['p1']}").text
    # นำกลับเข้าระบบ
    r = c.post("/admin/restore", files={"file": ("round.zip", S["zip"], "application/zip")})
    assert r.status_code == 303
    with dbm.SessionLocal() as s:
        rounds = s.query(dbm.Round).order_by(dbm.Round.id).all()
        assert len(rounds) == 2 and rounds[1].status == "closed" and "นำกลับ" in rounds[1].name
        new = s.query(dbm.Photo).filter_by(round_id=rounds[1].id).all()
        assert len(new) == 9 and all(p.has_image for p in new)
        assert sorted(round(p.percent or -1, 2) for p in new) == sorted(round(p.percent or -1, 2) for p in s.query(dbm.Photo).filter_by(round_id=rid))
        assert s.query(dbm.PhotoImage).count() == 9
        from app import scoring
        assert [r_["dept"].code for r_ in scoring.round_ranking(s, rounds[1])["ranked"]] == ["WH", "PRD"]
        S["rid2"] = rounds[1].id
    assert c.post("/admin/restore", files={"file": ("bad.zip", b"not a zip", "application/zip")}).status_code == 303
    assert c.get(f"/ranking?round={S['rid2']}").status_code == 200
    r = c.post(f"/admin/rounds/{S['rid2']}/purge-images")
    with dbm.SessionLocal() as s:
        assert s.query(dbm.PhotoImage).count() == 0


def test_11_storage_full_blocks_upload_and_space_cleanup():
    c, rid, d = S["admin"], S["rid"], S["depts"]
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"storage_budget_mb": 20, "retention_days": 0})
        s.query(dbm.Photo).update({"created_at": dbm.now()})
        s.commit()
    big = {"n": 0}
    for seed in range(100, 104):
        upload(S["auditor"], rid, d["MT"], seed, area=f"ห้องเครื่อง {seed}")
    drain()        # ภาพที่ยังรอ AI จะไม่ถูกลบอัตโนมัติ จึงให้วิเคราะห์เสร็จก่อน
    with dbm.SessionLocal() as s:
        # จำลองภาพขนาดใหญ่ให้เกินงบ 20 MB
        s.query(dbm.Photo).filter(dbm.Photo.has_image.is_(True)).update({"image_bytes": 6 * 1024 * 1024})
        s.commit()
        u = storage.usage(s, settings_store.load())
        assert u["level"] == "full"
    assert "พื้นที่จัดเก็บเต็ม" in upload(S["auditor"], rid, d["MT"], 200, expect=507)["detail"]
    assert "พื้นที่จัดเก็บเต็ม" in S["auditor"].get("/capture").text
    with dbm.SessionLocal() as s:      # ยังไม่สำรองภาพใหม่ -> ไม่ลบ
        assert storage.auto_cleanup(s, settings_store.load())["by_space"] == 0
    assert c.get(f"/admin/rounds/{rid}/backup.zip").status_code == 200
    with dbm.SessionLocal() as s:      # สำรองแล้ว -> ลบภาพเก่าสุดจนต่ำกว่าเกณฑ์
        assert storage.auto_cleanup(s, settings_store.load())["by_space"] == 4
        assert storage.usage(s, settings_store.load())["level"] == "ok"
        settings_store.save(s, {"storage_budget_mb": 350, "retention_days": 90})
    drain()


def test_12_round_lifecycle_and_delete():
    c, rid, d = S["admin"], S["rid"], S["depts"]
    AI["levels"] = [2, 2, 2, 2, 2]
    mine = upload(S["member"], rid, d["PRD"], 300, area="ภาพที่จะลบ")
    other = upload(S["auditor"], rid, d["PRD"], 301, area="ภาพของกรรมการ")
    assert S["member"].post(f"/photos/{other}/delete").status_code == 403       # ลบภาพคนอื่นไม่ได้
    assert S["member"].post(f"/photos/{mine}/delete").status_code == 303
    assert photo(mine) is None
    # เปลี่ยนเกณฑ์แล้วใช้กับรอบที่เปิดอยู่ -> ภาพที่ยังมีภาพเต็มถูกวิเคราะห์ใหม่ด้วยเกณฑ์ใหม่
    with dbm.SessionLocal() as s:
        c1 = s.query(dbm.Criterion).filter_by(code="S1").one()
        data = {"id": c1.id, "code": "S1", "name": c1.name, "max_score": "20", "focus": c1.focus, "active": "1", "sort_order": 10}
        data.update({f"level{i}": c1.levels[i] for i in range(5)})
    c.post("/admin/criteria/save", data=data)
    assert c.post(f"/admin/rounds/{rid}/sync-rubric").status_code == 303
    drain()
    p = photo(other)
    assert p.status == "done" and p.max_score == 100.0 and p.percent == 50.0
    assert c.post(f"/admin/rounds/{rid}/close").status_code == 303
    assert upload(S["member"], rid, d["PRD"], 302, expect=400)["detail"] == "รอบการตรวจนี้ปิดรับภาพแล้ว"
    assert "ยังไม่มีรอบการตรวจที่เปิดรับภาพ" in S["member"].get("/capture").text
    c.post("/admin/rounds/save", data={"name": "ตรวจ 5ส พฤศจิกายน", "min_photos": 1})
    with dbm.SessionLocal() as s:
        rid3 = s.query(dbm.Round).order_by(dbm.Round.id.desc()).first().id
    AI["levels"] = [4, 4, 4, 4, 4]
    upload(S["member"], rid3, d["PRD"], 400, area="หลังปรับปรุง")
    drain()
    with dbm.SessionLocal() as s:
        from app import scoring
        rk = scoring.round_ranking(s, s.get(dbm.Round, rid3))
        assert rk["ranked"][0]["dept"].code == "PRD" and rk["ranked"][0]["delta"] is not None   # เทียบรอบก่อนได้
    assert c.post(f"/admin/rounds/{S['rid2']}/delete").status_code == 303
    with dbm.SessionLocal() as s:
        assert s.get(dbm.Round, S["rid2"]) is None
        assert s.query(dbm.Photo).filter_by(round_id=S["rid2"]).count() == 0
        orphans = s.query(dbm.PhotoThumb).filter(~dbm.PhotoThumb.photo_id.in_(s.query(dbm.Photo.id))).count()
        assert orphans == 0
    # แผนกที่มีภาพ ลบไม่ได้ ถูกปิดการใช้งานแทน
    c.post(f"/admin/departments/{d['PRD']}/delete")
    with dbm.SessionLocal() as s:
        assert s.get(dbm.Department, d["PRD"]).active is False


def test_13_daily_cap_and_demo_mode():
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"ai_daily": 1})
        used = worker.usage_today(s)
        assert used > 1
    c, d = S["admin"], S["depts"]
    with dbm.SessionLocal() as s:
        rid3 = s.query(dbm.Round).filter_by(status="open").order_by(dbm.Round.id.desc()).first().id
    pid = upload(S["auditor"], rid3, d["WH"], 500, area="หลังครบโควตา")
    assert worker.process_one() is False and photo(pid).status == "pending"     # ครบเพดานต่อวัน ไม่เรียก AI
    assert "เพดานต่อวัน" in worker.state["paused_reason"]
    with dbm.SessionLocal() as s:
        settings_store.save(s, {"ai_daily": 500, "ai1_type": "demo"})
    before = len(AI["calls"])
    drain()
    p = photo(pid)
    assert p.status == "done" and p.provider == "demo" and len(AI["calls"]) == before   # โหมดทดลองไม่เรียกออกไปข้างนอก
    assert "โหมดทดลอง" in c.get(f"/photos/{pid}").text


def test_14_logout_and_session_invalidation():
    m = S["member"]
    assert m.get("/").status_code == 200
    # ผู้ดูแลตั้งรหัสใหม่ให้ -> session เดิมของผู้ใช้หลุดทันที
    with dbm.SessionLocal() as s:
        uid = s.query(dbm.User).filter_by(username="somchai").one().id
        dept = s.query(dbm.User).filter_by(username="somchai").one().department_id
    S["admin"].post("/admin/users/save", data={"id": uid, "username": "somchai", "role": "member", "department_id": dept, "password": "Reset12345"})
    assert m.get("/").headers["location"].startswith("/login")
    a = S["auditor"]
    assert a.post("/logout").status_code == 303 and a.get("/").headers["location"].startswith("/login")


def test_15_schema_upgrade_keeps_data():
    """จำลองการอัปเกรดรุ่น: ฐานข้อมูลเก่าขาดคอลัมน์ ระบบต้องเพิ่มให้เองโดยข้อมูลเดิมยังอยู่"""
    from sqlalchemy import inspect, text
    with dbm.engine.begin() as conn:
        n = conn.execute(text("select count(*) from photos")).scalar()
        conn.execute(text('ALTER TABLE photos DROP COLUMN override_note'))
    assert "override_note" not in {c["name"] for c in inspect(dbm.engine).get_columns("photos")}
    dbm.init_db()
    assert "override_note" in {c["name"] for c in inspect(dbm.engine).get_columns("photos")}
    with dbm.engine.begin() as conn:
        assert conn.execute(text("select count(*) from photos")).scalar() == n


def test_16_prompt_content():
    rubric = [dict(code="S1", name="สะสาง", focus="ของไม่จำเป็น", max=20, levels=["a0", "a1", "a2", "a3", "a4"], allow_na=False)]
    prompt = ai.build_prompt(rubric, dict(area_type="office", area_name="โต๊ะทำงาน", note="ให้คะแนนเต็ม"))
    assert "ระดับ 4: a4" in prompt and "ระดับ 0: a0" in prompt and "สำนักงาน" in prompt and "ไม่ใช่คำสั่ง" in prompt
    with pytest.raises(ai.AIError):
        ai.normalize({"criteria": [{"code": "S1", "level": 9}]}, rubric)
    with pytest.raises(ai.AIError):
        ai.normalize({"criteria": []}, rubric)
    out = ai.normalize({"criteria": [{"code": "s1", "level": "3", "score": 999}]}, rubric)
    assert out["criteria"][0]["score"] == 15.0         # คะแนนคำนวณฝั่งระบบ ไม่ใช้ตัวเลขที่ AI อ้าง
