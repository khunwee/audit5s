"""สำรองข้อมูล / นำกลับ / ส่งออก Excel / สร้างรายงาน

ไฟล์สำรองของรอบ (ZIP) เปิดดูได้โดยไม่ต้องมีระบบ: มี report.html, scores.xlsx, รูปทุกใบ
และ manifest.json ที่ใช้นำรอบนั้นกลับเข้าระบบได้ทั้งหมด
"""
import csv
import io
import json
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config, imaging, scoring, settings_store
from .photos import SOURCES, area_label
from .db import (Criterion, Department, DeptSummary, Photo, PhotoImage, PhotoThumb,
                 Round, User, log, now)
from .web import STATUS, f_date, f_dt, f_num, templates

FORMAT = "fives-round-backup/1"
_CSS = Path(__file__).parent / "static" / "app.css"


def _iso(v):
    return v.isoformat() if isinstance(v, (datetime, date)) else v


def _dt(v):
    try:
        return datetime.fromisoformat(v) if v else None
    except Exception:
        return None


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "_", name or "x")[:30] or "x"


PHOTO_FIELDS = ["area_name", "area_type", "note", "uploader_name", "sha256", "width", "height",
                "status", "error", "provider", "model", "score", "max_score", "percent", "analysis",
                "overridden", "override_by", "override_note", "source", "review_flag"]
PHOTO_DATES = ["created_at", "analyzed_at", "override_at", "purged_at"]


def photo_dict(p: Photo) -> dict:
    d = {k: getattr(p, k) for k in PHOTO_FIELDS}
    d.update({k: _iso(getattr(p, k)) for k in PHOTO_DATES})
    d.update(id=p.id, dept_code=p.department.code)
    return d


# --------------------------------------------------------------------------- รายงาน
def report_context(db, rnd: Round, img=None) -> dict:
    ranking = scoring.round_ranking(db, rnd)
    photos = db.query(Photo).filter(Photo.round_id == rnd.id).order_by(Photo.percent.desc(), Photo.id).all()
    by = defaultdict(list)
    for p in photos:
        by[p.department_id].append(p)
    sums = {x.department_id: x for x in db.query(DeptSummary).filter(DeptSummary.round_id == rnd.id)
            .order_by(DeptSummary.id).all()}
    s = settings_store.load()
    return dict(rnd=rnd, ranking=ranking, by=by, summaries=sums, org_name=s.get("org_name", ""),
                img=img or (lambda p: f"/photos/{p.id}/image"), generated=now(),
                css=_CSS.read_text(encoding="utf-8"))


def render_report(db, rnd: Round, img=None) -> str:
    return templates.env.get_template("report_full.html").render(**report_context(db, rnd, img))


# --------------------------------------------------------------------------- Excel
def export_excel(db, rnd: Round) -> bytes:
    ranking = scoring.round_ranking(db, rnd)
    rubric = ranking["rubric"]
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1D2A2F")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")

    def header(ws, row, titles):
        for i, t in enumerate(titles, 1):
            c = ws.cell(row=row, column=i, value=t)
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(wrap_text=True, vertical="center")

    ws = wb.active
    ws.title = "อันดับ"
    ws["A1"] = f"ผลการตรวจ 5ส: {rnd.name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"ช่วงเวลา {f_date(rnd.start_date)} ถึง {f_date(rnd.end_date)}   คะแนนเต็ม {f_num(ranking['full_score'])}"
                f"   ภาพขั้นต่ำต่อแผนก {rnd.min_photos}   ออกรายงาน {f_dt(now())}")
    titles = ["อันดับ", "รหัส", "แผนก", "คะแนนเฉลี่ย (%)", "ภาพที่ให้คะแนน", "ภาพทั้งหมด",
              "ภาพคะแนนต่ำสุด (%)", "ภาพคะแนนสูงสุด (%)", "เทียบรอบก่อน"] + [f"{c['name']} (%)" for c in rubric] + ["สถานะ"]
    header(ws, 4, titles)
    r = 5
    for group, label in (("ranked", "จัดอันดับแล้ว"), ("unranked", "ภาพยังไม่ครบ"), ("idle", "ยังไม่ส่งภาพ")):
        for row in ranking[group]:
            vals = [row["rank"], row["dept"].code, row["dept"].name, row["avg"], row["scored"], row["total"],
                    row["low"], row["high"], row["delta"]] + [row["crit"].get(c["code"]) for c in rubric] + [label]
            for i, v in enumerate(vals, 1):
                ws.cell(row=r, column=i, value=v)
            r += 1
    for i, w in enumerate([8, 10, 28, 16, 14, 12, 16, 16, 14] + [18] * len(rubric) + [16], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    ws2 = wb.create_sheet("รายละเอียดภาพ")
    titles = ["ภาพ", "เลขที่", "แผนก", "จุดตรวจ", "ประเภทพื้นที่", "ผู้ส่ง", "วันที่ส่ง", "สถานะ",
              "คะแนน", "คะแนนเต็ม", "%"]
    for c in rubric:
        titles += [f"{c['code']} ระดับ", f"{c['code']} คะแนน"]
    titles += ["สรุป", "เหตุผลรายเกณฑ์", "คำแนะนำ", "ปรับคะแนนโดย", "เหตุผลที่ปรับ", "โมเดล"]
    header(ws2, 1, titles)
    photos = (db.query(Photo).filter(Photo.round_id == rnd.id)
              .order_by(Photo.department_id, Photo.id).all())
    keep = []
    for n, p in enumerate(photos, start=2):
        a = p.analysis or {}
        crit = {c["code"]: c for c in a.get("criteria", [])}
        vals = ["", p.id, p.department.name, p.area_name, area_label(p.area_type),
                p.uploader_name, f_dt(p.created_at), STATUS.get(p.status, p.status),
                p.score, p.max_score, p.percent]
        for c in rubric:
            item = crit.get(c["code"])
            if item is None:
                vals += ["", ""]
            elif item.get("na"):
                vals += ["ประเมินไม่ได้", ""]
            else:
                vals += [item.get("level"), item.get("score")]
        reasons = "\n".join(f"{c['name']}: {c.get('reason', '')}" for c in a.get("criteria", []) if not c.get("na"))
        recs = "\n".join(f"- {x}" for c in a.get("criteria", []) for x in c.get("recommendations", []))
        vals += [a.get("summary") or a.get("image_issue") or p.error, reasons, recs,
                 p.override_by if p.overridden else "", p.override_note if p.overridden else "", p.model]
        for i, v in enumerate(vals, 1):
            ws2.cell(row=n, column=i, value=v).alignment = wrap
        ws2.row_dimensions[n].height = 62
        thumb = db.get(PhotoThumb, p.id)
        if thumb is not None:
            try:
                buf = io.BytesIO(thumb.data)
                pic = XLImage(buf)
                scale = 76 / max(pic.height, 1)
                pic.width, pic.height = int(pic.width * scale), 76
                ws2.add_image(pic, f"A{n}")
                keep.append(buf)
            except Exception:
                pass
            db.expunge(thumb)
    widths = [16, 8, 22, 22, 18, 16, 16, 18, 9, 10, 8] + [10, 10] * len(rubric) + [40, 60, 60, 16, 30, 22]
    for i, w in enumerate(widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "C2"

    ws3 = wb.create_sheet("เกณฑ์ของรอบนี้")
    header(ws3, 1, ["รหัส", "เกณฑ์", "คะแนนเต็ม", "สิ่งที่ต้องดู", "ระดับ 4", "ระดับ 3", "ระดับ 2", "ระดับ 1", "ระดับ 0", "ข้ามได้"])
    for n, c in enumerate(rubric, start=2):
        lv = (list(c.get("levels") or []) + [""] * 5)[:5]
        vals = [c["code"], c["name"], c["max"], c.get("focus", ""), lv[4], lv[3], lv[2], lv[1], lv[0],
                "ได้" if c.get("allow_na") else "ไม่ได้"]
        for i, v in enumerate(vals, 1):
            ws3.cell(row=n, column=i, value=v).alignment = wrap
    for i, w in enumerate([8, 26, 10, 44, 36, 36, 36, 36, 36, 8], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# --------------------------------------------------------------------------- CSV / JSON
def _csv(rows: list) -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerows(rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")       # BOM ให้ Excel อ่านภาษาไทยถูก


def ranking_csv(db, rnd: Round) -> bytes:
    rk = scoring.round_ranking(db, rnd)
    rows = [["อันดับ", "รหัสแผนก", "แผนก", "คะแนนเฉลี่ย (%)", "ภาพที่ให้คะแนน", "ภาพทั้งหมด", "ต่ำสุด (%)", "สูงสุด (%)",
             "เทียบรอบก่อน"] + [f"{c['name']} (%)" for c in rk["rubric"]] + ["สถานะ"]]
    for group, label in (("ranked", "จัดอันดับแล้ว"), ("unranked", "ภาพยังไม่ครบ"), ("idle", "ยังไม่ส่งภาพ")):
        for r in rk[group]:
            rows.append([r["rank"] or "", r["dept"].code, r["dept"].name, r["avg"], r["scored"], r["total"], r["low"],
                         r["high"], r["delta"]] + [r["crit"].get(c["code"]) for c in rk["rubric"]] + [label])
    return _csv([["" if v is None else v for v in row] for row in rows])


def photos_csv(db, rnd: Round, dept_id: int = 0) -> bytes:
    rubric = list(rnd.rubric or [])
    head = ["เลขที่ภาพ", "รหัสแผนก", "แผนก", "จุดตรวจ", "ประเภทพื้นที่", "แหล่งภาพ", "ผู้ส่ง", "วันที่ส่ง", "สถานะ", "คะแนน",
            "คะแนนเต็ม", "%"]
    for c in rubric:
        head += [f"{c['code']} ระดับ", f"{c['code']} คะแนน", f"{c['code']} เหตุผล"]
    head += ["สรุป", "ควรทำก่อน", "ปรับโดยกรรมการ", "เหตุผลที่ปรับ", "ควรให้กรรมการตรวจ", "ภาพหลังแก้ไขของภาพเลขที่", "โมเดล"]
    q = db.query(Photo).filter(Photo.round_id == rnd.id)
    if dept_id:
        q = q.filter(Photo.department_id == dept_id)
    rows = [head]
    for p in q.order_by(Photo.department_id, Photo.id).all():
        a = p.analysis or {}
        crit = {c["code"]: c for c in a.get("criteria", [])}
        row = [p.id, p.department.code, p.department.name, p.area_name, area_label(p.area_type),
               SOURCES.get(p.source or "mobile", p.source), p.uploader_name, f_dt(p.created_at), STATUS.get(p.status, p.status),
               p.score, p.max_score, p.percent]
        for c in rubric:
            item = crit.get(c["code"])
            if item is None:
                row += ["", "", ""]
            elif item.get("na"):
                row += ["ประเมินไม่ได้", "", item.get("reason", "")]
            else:
                row += [item.get("level"), item.get("score"), item.get("reason", "")]
        row += [a.get("summary") or a.get("image_issue") or p.error, " | ".join(a.get("top_actions") or []),
                p.override_by if p.overridden else "", p.override_note if p.overridden else "",
                "ใช่" if p.review_flag else "", p.after_of or "", p.model]
        rows.append(["" if v is None else v for v in row])
    return _csv(rows)


def round_json(db, rnd: Round) -> bytes:
    rk = scoring.round_ranking(db, rnd)

    def row(r):
        return dict(rank=r["rank"], code=r["dept"].code, department=r["dept"].name, average=r["avg"], scored=r["scored"],
                    total=r["total"], low=r["low"], high=r["high"], delta=r["delta"], criteria=r["crit"])
    photos = db.query(Photo).filter(Photo.round_id == rnd.id).order_by(Photo.department_id, Photo.id).all()
    data = dict(format="fives-round-export/1", exported_at=_iso(now()),
                round=dict(id=rnd.id, name=rnd.name, status=rnd.status, start_date=_iso(rnd.start_date),
                           end_date=_iso(rnd.end_date), min_photos=rnd.min_photos, full_score=rk["full_score"],
                           rubric=rnd.rubric),
                ranking=[row(r) for r in rk["ranked"]], not_ranked=[row(r) for r in rk["unranked"]],
                no_photos=[r["dept"].name for r in rk["idle"]], photos=[photo_dict(p) for p in photos])
    return json.dumps(data, ensure_ascii=False, indent=1).encode("utf-8")


# --------------------------------------------------------------------------- สำรองรอบ
def build_round_zip(db, rnd: Round, path: str, user=None) -> dict:
    photos = db.query(Photo).filter(Photo.round_id == rnd.id).order_by(Photo.department_id, Photo.id).all()
    files, n_full = {}, 0
    with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as z:
        for p in photos:
            blob = db.get(PhotoImage, p.id) if p.has_image else None
            if blob is not None:
                name = f"images/{_safe(p.department.code)}/{p.id}.jpg"
                z.writestr(name, blob.data)
                db.expunge(blob)
                n_full += 1
            else:
                thumb = db.get(PhotoThumb, p.id)
                if thumb is None:
                    continue
                name = f"thumbs/{p.id}.jpg"
                z.writestr(name, thumb.data)
                db.expunge(thumb)
            files[p.id] = name
        depts = {p.department_id: p.department for p in photos}
        manifest = dict(
            format=FORMAT, app_version=config.APP_VERSION, exported_at=_iso(now()),
            round=dict(name=rnd.name, note=rnd.note, start_date=_iso(rnd.start_date), end_date=_iso(rnd.end_date),
                       status=rnd.status, min_photos=rnd.min_photos, rubric=rnd.rubric,
                       created_at=_iso(rnd.created_at), closed_at=_iso(rnd.closed_at)),
            departments=[dict(code=d.code, name=d.name, zone=d.zone) for d in depts.values()],
            photos=[dict(photo_dict(p), file=files.get(p.id)) for p in photos],
            summaries=[dict(dept_code=depts[x.department_id].code if x.department_id in depts else None,
                            content=x.content, model=x.model, created_at=_iso(x.created_at))
                       for x in db.query(DeptSummary).filter(DeptSummary.round_id == rnd.id).all()])
        z.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=1),
                   compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("report.html", render_report(db, rnd, img=lambda p: files.get(p.id, "")),
                   compress_type=zipfile.ZIP_DEFLATED)
        z.writestr("scores.xlsx", export_excel(db, rnd))
        z.writestr("อ่านก่อน.txt",
                   "ไฟล์สำรองผลการตรวจ 5ส\n"
                   "- report.html  เปิดด้วยเบราว์เซอร์เพื่อดูรายงานพร้อมรูป (ต้องแตกไฟล์ ZIP ก่อน)\n"
                   "- scores.xlsx  คะแนนและรายละเอียดทุกภาพ\n"
                   "- images/      ภาพเต็ม แยกโฟลเดอร์ตามรหัสแผนก\n"
                   "- manifest.json ใช้นำรอบนี้กลับเข้าระบบที่หน้า พื้นที่และสำรองข้อมูล\n")
    rnd.last_backup_at = now()
    log(db, user, "backup_round", f"{rnd.name}: {len(photos)} ภาพ (ภาพเต็ม {n_full})")
    db.commit()
    return dict(photos=len(photos), full=n_full)


def restore_round_zip(db, fileobj, user=None) -> Round:
    try:
        z = zipfile.ZipFile(fileobj)
        manifest = json.loads(z.read("manifest.json").decode("utf-8"))
    except Exception:
        raise ValueError("ไฟล์นี้ไม่ใช่ไฟล์สำรองของรอบการตรวจ (ไม่พบ manifest.json)")
    if manifest.get("format") != FORMAT:
        raise ValueError("รูปแบบไฟล์สำรองไม่ตรงกับระบบรุ่นนี้")
    names = {i.filename: i for i in z.infolist()}
    r = manifest["round"]
    dmap = {d.code: d for d in db.query(Department).all()}
    for d in manifest.get("departments", []):
        if d["code"] not in dmap:
            dep = Department(code=d["code"], name=d.get("name") or d["code"], zone=d.get("zone") or "")
            db.add(dep)
            dmap[d["code"]] = dep
    db.flush()
    name = r["name"]
    if db.query(Round).filter(Round.name == name).first():
        name = f"{name} (นำกลับ {f_dt(now())})"
    rnd = Round(name=name[:160], note=r.get("note") or "", status="closed", min_photos=int(r.get("min_photos") or 1),
                rubric=r.get("rubric") or [], closed_at=_dt(r.get("closed_at")) or now(),
                created_at=_dt(r.get("created_at")) or now(), last_backup_at=now(),
                start_date=_dt(r.get("start_date")).date() if _dt(r.get("start_date")) else None,
                end_date=_dt(r.get("end_date")).date() if _dt(r.get("end_date")) else None)
    db.add(rnd)
    db.flush()
    count = 0
    for item in manifest.get("photos", []):
        dep = dmap.get(item.get("dept_code"))
        if dep is None:
            continue
        p = Photo(round_id=rnd.id, department_id=dep.id, uploader_id=None, has_image=False)
        for k in PHOTO_FIELDS:
            if k in item:
                setattr(p, k, item[k])
        for k in PHOTO_DATES:
            setattr(p, k, _dt(item.get(k)))
        p.created_at = p.created_at or now()
        if p.status in ("pending", "processing"):
            p.status = "error"
            p.error = "ภาพนี้ยังไม่ได้วิเคราะห์ตอนสำรองข้อมูล กดวิเคราะห์ใหม่ได้"
        db.add(p)
        db.flush()
        info = names.get(item.get("file") or "")
        if info is not None and info.file_size <= 25 * 1024 * 1024:
            data = z.read(info)
            if info.filename.startswith("images/"):
                try:
                    out = imaging.process(data, 2400, 90)
                    db.add(PhotoImage(photo_id=p.id, data=data))
                    db.add(PhotoThumb(photo_id=p.id, data=out["thumb"]))
                    p.has_image, p.purged_at = True, None
                    p.image_bytes, p.thumb_bytes = len(data), len(out["thumb"])
                except imaging.ImageError:
                    pass
            else:
                db.add(PhotoThumb(photo_id=p.id, data=data))
                p.thumb_bytes = len(data)
        count += 1
        if count % 20 == 0:
            db.commit()
    for x in manifest.get("summaries", []):
        dep = dmap.get(x.get("dept_code"))
        if dep is not None:
            db.add(DeptSummary(round_id=rnd.id, department_id=dep.id, content=x.get("content"),
                               model=x.get("model") or "", created_at=_dt(x.get("created_at")) or now()))
    log(db, user, "restore_round", f"{rnd.name}: {count} ภาพ")
    db.commit()
    return rnd


# --------------------------------------------------------------------------- สำรองข้อมูลระบบ
def system_json(db) -> dict:
    """ข้อมูลทั้งระบบแบบข้อความ (ไม่มีรูป ไม่มี API key) สำหรับเก็บไว้อ้างอิงหรือย้ายระบบ"""
    s = {k: v for k, v in settings_store.load().items() if not k.endswith("_key") and not k.startswith("_")}
    depts = {d.id: d for d in db.query(Department).all()}
    return dict(
        format="fives-system-backup/1", app_version=config.APP_VERSION, exported_at=_iso(now()), settings=s,
        departments=[dict(code=d.code, name=d.name, zone=d.zone, active=d.active) for d in depts.values()],
        users=[dict(username=u.username, full_name=u.full_name, role=u.role, active=u.active,
                    department=u.department.code if u.department else None) for u in db.query(User).all()],
        criteria=[dict(code=c.code, name=c.name, focus=c.focus, max_score=c.max_score, levels=c.levels,
                       allow_na=c.allow_na, sort_order=c.sort_order, active=c.active)
                  for c in db.query(Criterion).order_by(Criterion.sort_order).all()],
        rounds=[dict(id=r.id, name=r.name, note=r.note, status=r.status, min_photos=r.min_photos,
                     start_date=_iso(r.start_date), end_date=_iso(r.end_date), rubric=r.rubric,
                     last_backup_at=_iso(r.last_backup_at),
                     photos=[photo_dict(p) for p in db.query(Photo).filter(Photo.round_id == r.id).order_by(Photo.id)])
                for r in db.query(Round).order_by(Round.id).all()])
